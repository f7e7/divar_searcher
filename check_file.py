#!/usr/bin/env python3
import openpyxl
import os

def check_excel_file():
    filepath = 'excel_files/divar_search_results_سمند_سورن_پلاس_20251231_081143.xlsx'

    if not os.path.exists(filepath):
        print(f'File does not exist: {filepath}')
        return

    try:
        wb = openpyxl.load_workbook(filepath)
        ws = wb.active

        print(f'File: {filepath}')
        print(f'Rows: {ws.max_row}, Columns: {ws.max_column}')

        headers = [ws.cell(row=1, column=col).value for col in range(1, ws.max_column + 1)]
        print(f'Headers: {headers}')

        if ws.max_row > 1:
            print('Data rows:')
            for row in range(2, min(6, ws.max_row + 1)):
                title = ws.cell(row=row, column=1).value
                phone = ws.cell(row=row, column=4).value
                print(f'  Row {row-1}: Title={title}, Phone={phone}')
        else:
            print('NO DATA ROWS - FILE IS EMPTY!')

    except Exception as e:
        print(f'Error: {e}')

if __name__ == '__main__':
    check_excel_file()
