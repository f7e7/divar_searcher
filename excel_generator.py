import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
import os
import time
import config

class ExcelGenerator:
    def __init__(self):
        self.workbook = None
        self.worksheet = None

    def create_excel(self, search_term, data):
        """
        Create Excel file with search results
        """
        timestamp = time.strftime("%Y%m%d_%H%M%S")
        filename = config.EXCEL_FILENAME_TEMPLATE.format(
            search_term=search_term.replace(' ', '_'),
            timestamp=timestamp
        )

        # Filter and clean data
        filtered_data = self._filter_and_clean_data(data)

        # If no data after filtering, return None
        if not filtered_data:
            print("No valid data found after filtering")
            return None

        # Create workbook and worksheet
        self.workbook = openpyxl.Workbook()
        self.worksheet = self.workbook.active
        self.worksheet.title = "Divar Search Results"

        # Add headers
        headers = ['Title', 'Description', 'Price', 'URL']
        self._add_headers(headers)

        # Add data
        self._add_data(filtered_data)

        # Auto-adjust column widths
        self._adjust_column_widths()

        # Ensure excel_files directory exists
        excel_dir = "excel_files"
        os.makedirs(excel_dir, exist_ok=True)

        # Save file to excel_files directory
        filepath = os.path.join(excel_dir, filename)
        self.workbook.save(filepath)

        print(f"Excel file created: {filepath} with {len(filtered_data)} records")
        return filepath

    def _filter_and_clean_data(self, data):
        """
        Filter out invalid entries and remove duplicates
        """
        filtered_data = []
        seen_urls = set()

        for item in data:
            title = item.get('title', '').strip()
            url = item.get('url', '').strip()

            # Skip items without title or with placeholder titles
            if not title or title.lower() in ['no title', 'بدون عنوان', '']:
                continue

            # Skip items without URL
            if not url:
                continue

            # Skip duplicates based on URL (more reliable than title)
            if url in seen_urls:
                continue

            # Add to filtered data and mark as seen
            filtered_data.append(item)
            seen_urls.add(url)

        return filtered_data

    def _add_headers(self, headers):
        """
        Add headers to the worksheet with styling
        """
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        alignment = Alignment(horizontal="center", vertical="center")

        for col_num, header in enumerate(headers, 1):
            cell = self.worksheet.cell(row=1, column=col_num)
            cell.value = header
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = alignment

    def _add_data(self, data):
        """
        Add scraped data to the worksheet
        """
        alignment = Alignment(wrap_text=True, vertical="top")

        for row_num, item in enumerate(data, 2):
            # Title
            cell = self.worksheet.cell(row=row_num, column=1)
            cell.value = item.get('title', '')
            cell.alignment = alignment

            # Description
            cell = self.worksheet.cell(row=row_num, column=2)
            cell.value = item.get('description', '')
            cell.alignment = alignment

            # Price
            cell = self.worksheet.cell(row=row_num, column=3)
            cell.value = item.get('price', '')
            cell.alignment = alignment

            # URL
            cell = self.worksheet.cell(row=row_num, column=4)
            cell.value = item.get('url', '')
            cell.alignment = alignment

    def _adjust_column_widths(self):
        """
        Auto-adjust column widths based on content
        """
        for column in self.worksheet.columns:
            max_length = 0
            column_letter = column[0].column_letter

            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass

            # Set minimum width of 10, maximum of 50
            adjusted_width = min(max(max_length + 2, 10), 50)
            self.worksheet.column_dimensions[column_letter].width = adjusted_width

    def cleanup_old_files(self, keep_days=7):
        """
        Clean up old Excel files to save disk space
        """
        current_time = time.time()
        cutoff_time = current_time - (keep_days * 24 * 60 * 60)

        for filename in os.listdir('.'):
            if filename.startswith('divar_search_results_') and filename.endswith('.xlsx'):
                filepath = os.path.join('.', filename)
                if os.path.getmtime(filepath) < cutoff_time:
                    try:
                        os.remove(filepath)
                        print(f"Cleaned up old file: {filename}")
                    except Exception as e:
                        print(f"Error removing {filename}: {e}")





