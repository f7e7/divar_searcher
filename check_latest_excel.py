#!/usr/bin/env python3
import openpyxl
import os

def check_latest_excel():
    excel_dir = 'excel_files'
    files = [f for f in os.listdir(excel_dir) if f.endswith('.xlsx')]
    if not files:
        print('No Excel files found')
        return

    latest_file = max([os.path.join(excel_dir, f) for f in files], key=os.path.getmtime)
    print(f'Checking: {os.path.basename(latest_file)}')

    wb = openpyxl.load_workbook(latest_file)
    ws = wb.active

    print(f'Rows: {ws.max_row}, Columns: {ws.max_column}')
    print(f'Data rows: {ws.max_row - 1}')

    for row in range(2, min(6, ws.max_row + 1)):
        title = ws.cell(row=row, column=1).value
        phone = ws.cell(row=row, column=4).value
        print(f'{row-1}. {title} | Phone: {phone}')

if __name__ == '__main__':
    check_latest_excel()
